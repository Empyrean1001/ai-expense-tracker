from decimal import Decimal

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.dashboard import SHEET_ORDER, generate_dashboard
from src.insights import (
    compute_month_over_month_deltas,
    compute_top_vendors,
    detect_recurring_charges,
)
from src.transactions_store import CSV_COLUMNS, append_transactions, read_all_transactions
from tests.test_categorizer import make_txn


def history_df():
    rows = [
        dict(source_month="2026-05", vendor="Netflix", category="Entertainment",
             subcategory="Streaming", amount=199.0, direction="debit"),
        dict(source_month="2026-06", vendor="Netflix", category="Entertainment",
             subcategory="Streaming", amount=199.0, direction="debit"),
        dict(source_month="2026-07", vendor="Netflix", category="Entertainment",
             subcategory="Streaming", amount=199.0, direction="debit"),
        dict(source_month="2026-06", vendor="Swiggy", category="Food & Dining",
             subcategory="Delivery", amount=1000.0, direction="debit"),
        dict(source_month="2026-07", vendor="Swiggy", category="Food & Dining",
             subcategory="Delivery", amount=1350.0, direction="debit"),
        dict(source_month="2026-07", vendor="ACME", category="Income",
             subcategory="Salary", amount=50000.0, direction="credit"),
    ]
    return pd.DataFrame(rows)


class TestInsights:
    def test_top_vendors_ignores_credits(self):
        top = compute_top_vendors(history_df(), "2026-07")
        assert list(top["vendor"]) == ["Swiggy", "Netflix"]

    def test_recurring_detects_netflix_only(self):
        recurring = detect_recurring_charges(history_df())
        assert list(recurring["vendor"]) == ["Netflix"]
        assert recurring.iloc[0]["occurrences"] == 3

    def test_mom_deltas(self):
        deltas = compute_month_over_month_deltas(history_df())
        food_july = deltas[
            (deltas["category"] == "Food & Dining") & (deltas["source_month"] == "2026-07")
        ].iloc[0]
        assert food_july["pct_change"] == pytest.approx(35.0)
        food_june = deltas[
            (deltas["category"] == "Food & Dining") & (deltas["source_month"] == "2026-06")
        ].iloc[0]
        assert pd.isna(food_june["pct_change"])


def store_df(tmp_path, txns):
    path = tmp_path / "t.csv"
    append_transactions(path, txns)
    return read_all_transactions(path)


class TestDashboard:
    def test_single_currency_has_six_base_sheets(self, tmp_path):
        df = store_df(tmp_path, [
            make_txn("SWIGGY*1", vendor="Swiggy", category="Food & Dining",
                     subcategory="Delivery", amount=Decimal("249.00")),
            make_txn("NEFT SALARY", vendor="ACME", category="Income",
                     subcategory="Salary", amount=Decimal("1500.00"), direction="credit"),
        ])
        out = generate_dashboard(df, [], tmp_path / "dash.xlsx", "2026-07")
        wb = load_workbook(out)
        assert wb.sheetnames == SHEET_ORDER

    def test_two_currencies_get_suffixed_sheets(self, tmp_path):
        df = store_df(tmp_path, [
            make_txn("SWIGGY*1", vendor="Swiggy", category="Food & Dining",
                     amount=Decimal("249.00"), currency="INR"),
            make_txn("EDEKA", vendor="EDEKA", category="Food & Dining",
                     amount=Decimal("45.67"), currency="EUR"),
        ])
        out = generate_dashboard(df, [], tmp_path / "dash.xlsx", "2026-07")
        wb = load_workbook(out)
        assert "Dashboard EUR" in wb.sheetnames
        assert "Dashboard INR" in wb.sheetnames
        assert len(wb.sheetnames) == 12

    def test_empty_history_still_writes_valid_workbook(self, tmp_path):
        df = pd.DataFrame(columns=CSV_COLUMNS)
        out = generate_dashboard(df, [], tmp_path / "dash.xlsx", "2026-07")
        wb = load_workbook(out)  # would raise if corrupt
        assert len(wb.sheetnames) == 6

    def test_ledger_summary_appears(self, tmp_path):
        df = store_df(tmp_path, [make_txn("X", vendor="X", amount=Decimal("10"))])
        ledger = [{
            "type": "direct", "person": "Rahul", "direction": "lent",
            "amount": "500", "currency": "INR", "entry_date": "2026-07-12",
            "status": "unsettled", "linked_transaction_hash": "h1" * 20,
        }]
        out = generate_dashboard(df, ledger, tmp_path / "dash.xlsx", "2026-07")
        ws = load_workbook(out)["Lend & Borrow Ledger"]
        assert "Owed to you: 500" in ws["A1"].value
