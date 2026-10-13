"""Ingestion orchestrator: any supported statement file in, canonical
transactions out.

Standard formats go straight to their deterministic parser. CSV/XLSX go
through the BankProfile machinery: cached profile if we've seen this bank
before; otherwise detect the layout from the file itself, then (only if that
fails) let the LLM label the header cells, then (only if that fails) map
the columns with the user — validate, ask anything the data can't prove,
cache. `model_factory` is lazy, so a run that never needs the LLM never even
constructs a client.
"""

import csv
from pathlib import Path
from typing import Callable

import click

from src.format_detector import detect_format
from src.layout_detect import detect_profile
from src.localeparse import disambiguate_date_order
from src.models import BankProfile, Transaction
from src.ollama_client import OllamaError
from src.profile import find_header_row, parse_rows_with_profile
from src.profile_store import match_profile, save_profile
from src.profile_validator import validate_profile
from src.schema_agent import (
    SchemaInferenceError,
    infer_profile,
    interactive_column_mapping,
)
from src.standard_formats import parse_camt053, parse_mt940, parse_ofx, parse_qif

SAMPLE_ROWS = 40


def read_raw_rows(path: Path, format_tag: str) -> list[list[str]]:
    if format_tag == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return [
                ["" if cell is None else str(cell) for cell in row]
                for row in workbook.active.iter_rows(values_only=True)
            ]
        finally:
            workbook.close()
    with open(path, newline="", encoding="utf-8-sig") as f:
        sample = f.read(4096)
        f.seek(0)
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        return list(csv.reader(f, delimiter=delimiter))


def _swap_day_month(fmt: str) -> str:
    return fmt.replace("%d", "\x00").replace("%m", "%d").replace("\x00", "%m")


def _resolve_ambiguities(
    rows: list[list[str]], profile: BankProfile, prompt_fn, default_currency: str
) -> BankProfile:
    """Ask the user the questions the data cannot answer — once, then the
    answers live in the cached profile forever."""
    updates: dict = {}

    date_col = profile.column_map.get("posted_at", profile.column_map.get("value_date"))
    header = find_header_row(rows, profile.header_signature)
    if header is not None and date_col is not None:
        samples = [
            row[date_col] for row in rows[header + 1 : header + 1 + SAMPLE_ROWS]
            if date_col < len(row) and row[date_col].strip()
        ]
        if disambiguate_date_order(samples) == "ambiguous" and any(
            "%d" in fmt and "%m" in fmt for fmt in profile.date_formats
        ):
            choice = str(prompt_fn(
                "This bank's dates are ambiguous (e.g. 03/05). Are they "
                "[d]ay-first or [m]onth-first?",
                type=click.Choice(["d", "m"]),
            )).strip().lower()
            day_first_in_profile = all(
                fmt.index("%d") < fmt.index("%m")
                for fmt in profile.date_formats
                if "%d" in fmt and "%m" in fmt
            )
            wants_day_first = choice == "d"
            if wants_day_first != day_first_in_profile:
                updates["date_formats"] = [_swap_day_month(f) for f in profile.date_formats]

    if not profile.currency:
        answer = str(prompt_fn(
            f"Which ISO currency is this account in? [{default_currency}]"
        )).strip().upper() or default_currency
        updates["currency"] = answer

    return profile.model_copy(update=updates) if updates else profile


def ingest(
    path: Path,
    *,
    profiles_dir: Path,
    model_factory: Callable[[], object],
    prompt_fn=click.prompt,
    echo_fn=click.echo,
    default_currency: str = "USD",
) -> list[Transaction]:
    format_tag = detect_format(path)

    if format_tag == "ofx":
        return parse_ofx(path)
    if format_tag == "qif":
        return parse_qif(path, currency=default_currency)
    if format_tag == "mt940":
        return parse_mt940(path)
    if format_tag == "camt053":
        return parse_camt053(path)

    rows = read_raw_rows(path, format_tag)

    cached = match_profile(profiles_dir, rows)
    if cached is not None:
        echo_fn(f"Recognized format: {cached.bank_name or 'saved bank profile'}.")
        return parse_rows_with_profile(rows, cached)

    echo_fn("New statement format — working out the layout (one-time step)...")
    profile = detect_profile(rows)
    if profile is None:
        try:
            profile = infer_profile(rows, model_factory())
        except (SchemaInferenceError, OllamaError) as exc:
            echo_fn(str(exc))
            profile = interactive_column_mapping(rows, prompt_fn, echo_fn)

    profile = _resolve_ambiguities(rows, profile, prompt_fn, default_currency)

    errors = validate_profile(rows, profile)
    if errors:
        echo_fn("The mapped layout still fails validation:\n  " + "\n  ".join(errors))
        profile = interactive_column_mapping(rows, prompt_fn, echo_fn)
        profile = _resolve_ambiguities(rows, profile, prompt_fn, default_currency)
        errors = validate_profile(rows, profile)
        if errors:
            raise SchemaInferenceError(
                "Manual mapping also failed validation — refusing to import "
                "possibly-misparsed money:\n  " + "\n  ".join(errors)
            )

    save_profile(profiles_dir, profile)
    echo_fn(f"Layout learned and saved — future statements from this bank parse instantly.")
    return parse_rows_with_profile(rows, profile)
